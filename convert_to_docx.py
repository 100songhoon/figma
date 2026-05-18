#!/usr/bin/env python3
"""Markdown to DOCX converter for the report."""
import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def main():
    doc = Document()
    
    # Page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # Read markdown
    with open('취준플랜_수집자료_보고서.md', 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_table = False
    table_rows = []
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                p = doc.add_paragraph()
                p.style = doc.styles['Normal']
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Courier New'
                run.font.size = Pt(8)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table handling
        if '|' in line and line.strip().startswith('|'):
            cells = [c.strip() for c in line.strip().split('|')[1:-1]]
            if cells and all(re.match(r'^[-:]+$', c) for c in cells):
                i += 1
                continue
            table_rows.append(cells)
            if not in_table:
                in_table = True
            i += 1
            continue
        elif in_table:
            # Flush table
            if table_rows:
                add_table(doc, table_rows)
                doc.add_paragraph()
            table_rows = []
            in_table = False

        # Skip horizontal rules
        if line.strip() == '---':
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.strip().startswith('> '):
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(line.strip()[2:])
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)
        elif line.strip().startswith('- [ ]') or line.strip().startswith('- [x]'):
            text = line.strip()[5:].strip()
            p = doc.add_paragraph(text, style='List Bullet')
        elif line.strip().startswith('- '):
            text = line.strip()[2:]
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Bullet')
        elif line.strip().startswith(re.findall(r'^\d+\.', line.strip())[0] if re.findall(r'^\d+\.', line.strip()) else 'NOMATCH'):
            text = re.sub(r'^\d+\.\s*', '', line.strip())
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Number')
        elif line.strip():
            text = line.strip()
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)
            doc.add_paragraph(text)

        i += 1

    # Flush remaining table
    if table_rows:
        add_table(doc, table_rows)

    doc.save('취준플랜_수집자료_보고서.docx')
    print("DOCX 파일 생성 완료: 취준플랜_수집자료_보고서.docx")


def add_table(doc, rows):
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            if ci < num_cols:
                cell = table.cell(ri, ci)
                cell_text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)
                cell_text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', cell_text)
                cell.text = cell_text
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    
    # Bold first row (header)
    if rows:
        for ci in range(num_cols):
            cell = table.cell(0, ci)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True


if __name__ == '__main__':
    main()
